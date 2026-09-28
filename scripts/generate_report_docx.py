"""Generate a concise 3-4 page report.docx from the assembled report_package/ assets.

Run scripts/generate_report_assets.py first to (re)build report_package/.
"""
from __future__ import annotations

import csv
import json
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[1]
PACKAGE_DIR = ROOT / "report_package"
OUTPUT_PATH = PACKAGE_DIR / "report.docx"


def add_title_page(doc: Document) -> None:
    title = doc.add_heading(
        "Pneumonia Opacity Detection: Pretrained Faster R-CNN with DINOv2 Semantic Fusion",
        level=0,
    )
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("RSNA Pneumonia Detection Challenge \u2014 Computer Vision Assignment")
    run.italic = True
    run.font.size = Pt(13)

    doc.add_paragraph()


def add_objective_and_dataset(doc: Document) -> None:
    doc.add_heading("1. Objective", level=1)
    doc.add_paragraph(
        "This project implements a computer-vision object-detection system that localizes "
        "pneumonia-related opacities in chest X-rays, preceded by a pretrained model as required "
        "by the assignment. Two phases were implemented: (1) a mandatory Faster R-CNN "
        "ResNet-50-FPN v2 baseline, pretrained on COCO/ImageNet, and (2) an enhancement that fuses "
        "a frozen DINOv2 ViT-S/14 global-context encoder into the detector at the ROI level."
    )

    doc.add_heading("2. Dataset", level=1)
    doc.add_paragraph(
        "Source: RSNA Pneumonia Detection Challenge (Kaggle/RSNA), chest X-rays in DICOM format. "
        "The full pool contains 26,684 training images and 30,227 labeled bounding-box rows. "
        "To fit the project timeline, all reported experiments use a stratified 12,000-patient "
        "subset, split 70/15/15 by patientId and positive-box presence (seed 42): 8,400 train / "
        "1,799 validation / 1,801 test images."
    )
    doc.add_paragraph(
        "Bounding boxes are converted once from the raw CSV format (x, y, width, height) to "
        "Torchvision's (x_min, y_min, x_max, y_max) convention. Multiple boxes for the same "
        "patient are merged into one multi-box training sample rather than treated as independent "
        "images."
    )


def add_architecture(doc: Document) -> None:
    doc.add_heading("3. Architecture", level=1)

    doc.add_heading("3.1 SPEC-1 \u2014 Baseline Detector", level=2)
    doc.add_paragraph(
        "DICOM X-ray \u2192 percentile normalization to [0,1] \u2192 replicate to 3 channels \u2192 "
        "pretrained ResNet-50 backbone \u2192 Feature Pyramid Network \u2192 Region Proposal Network \u2192 "
        "ROI Align \u2192 box classifier/regressor (re-initialized for 2 classes: background, "
        "pneumonia-opacity)."
    )

    doc.add_heading("3.2 SPEC-2 \u2014 DINOv2 Semantic Fusion", level=2)
    doc.add_paragraph(
        "The same SPEC-1 detector, up to ROI Align. A frozen DINOv2 ViT-S/14 encoder "
        "(requires_grad=False, always in .eval() mode) produces a 384-d global embedding, "
        "projected to 256-d via a trainable Linear+ReLU+Dropout layer. This projected embedding "
        "is broadcast and concatenated onto every per-proposal ROI feature, then passed through a "
        "fusion MLP before the final classification and box-regression heads. Only the fusion "
        "predictor and DINOv2 projection layer are trainable; the ResNet-50-FPN backbone/RPN/box_head "
        "and the DINOv2 backbone itself are frozen (initialized from the SPEC-1 baseline checkpoint)."
    )


def add_training_configuration_table(doc: Document) -> None:
    doc.add_heading("4. Training Configuration", level=1)
    table = doc.add_table(rows=1, cols=6)
    table.style = "Light Grid Accent 1"
    header = table.rows[0].cells
    for i, text in enumerate(["Run", "Pretrained", "Optimizer", "LR", "Scheduler", "Epochs"]):
        header[i].text = text

    rows = [
        ["Random-init (lower bound)", "No", "SGD", "0.003", "Cosine, warmup 300", "15"],
        ["Pretrained Baseline (SPEC-1)", "Yes (COCO)", "SGD", "0.003", "Cosine, warmup 300", "15"],
        ["DINOv2 Fusion (SPEC-2)", "Yes (COCO + DINOv2)", "AdamW", "0.001", "Cosine, warmup 200", "15"],
    ]
    for row_values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row_values):
            cells[i].text = value
    doc.add_paragraph()
    doc.add_paragraph(
        "All runs used the same 12,000-patient data split, batch size 8, weight decay 0.0005, "
        "gradient clipping (max norm 5.0), mixed precision, and image resize min_size=512/max_size=768."
    )


def add_image_with_caption(doc: Document, image_path: Path, caption: str, width_in: float = 6.0) -> None:
    if not image_path.exists():
        return
    doc.add_picture(str(image_path), width=Inches(width_in))
    caption_paragraph = doc.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = caption_paragraph.add_run(caption)
    run.italic = True
    run.font.size = Pt(9)


def add_training_graphs(doc: Document) -> None:
    doc.add_heading("5. Training Graphs", level=1)
    add_image_with_caption(
        doc,
        PACKAGE_DIR / "training_graphs" / "validation_map_overlay.png",
        "Figure 1. Validation mAP@0.50 across epochs for all three runs. The pretrained baseline "
        "peaks early then overfits; the DINOv2 fusion run stays stable; the random-init run is "
        "still improving at epoch 14.",
        width_in=6.0,
    )


def add_comparison_graphs(doc: Document) -> None:
    doc.add_heading("6. Comparison Graphs and Results", level=1)
    add_image_with_caption(
        doc,
        PACKAGE_DIR / "comparison_graphs" / "map_primary_metrics_comparison.png",
        "Figure 2. Primary detection metrics (mAP@0.50, mAP@0.50:0.95) across all three models.",
        width_in=5.5,
    )
    add_image_with_caption(
        doc,
        PACKAGE_DIR / "comparison_graphs" / "three_way_metric_comparison.png",
        "Figure 3. Full six-metric comparison across all three models.",
        width_in=6.0,
    )


def add_metrics_table(doc: Document) -> None:
    doc.add_heading("7. Ablation Metrics Table", level=1)
    csv_path = PACKAGE_DIR / "metrics" / "metrics_table.csv"
    with open(csv_path, newline="") as f:
        rows = list(csv.reader(f))

    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.style = "Light Grid Accent 1"
    for i, text in enumerate(rows[0]):
        table.rows[0].cells[i].text = text
    for data_row in rows[1:]:
        cells = table.add_row().cells
        for i, value in enumerate(data_row):
            try:
                cells[i].text = f"{float(value):.4f}"
            except ValueError:
                cells[i].text = value


def add_qualitative_predictions(doc: Document) -> None:
    doc.add_heading("8. Qualitative Predictions", level=1)
    doc.add_paragraph(
        "Ground truth boxes are drawn in green; predicted boxes are drawn in red with a "
        "confidence score."
    )
    sample_image = PACKAGE_DIR / "prediction_images" / "dinov2_fusion" / "prediction_03.png"
    add_image_with_caption(
        doc,
        sample_image,
        "Figure 4. DINOv2 fusion model prediction: correctly localized opacity (left lung) "
        "alongside a missed detection (right lung) \u2014 a representative true-positive / "
        "false-negative pair.",
        width_in=5.0,
    )


def add_conclusion(doc: Document) -> None:
    doc.add_heading("9. Conclusion", level=1)
    doc.add_paragraph(
        "1. Pretraining clearly improves detection performance: the pretrained Faster R-CNN "
        "baseline reaches 2.5x the mAP@0.50 of the random-init model (0.4673 vs 0.1847) under "
        "identical data, epochs, and hyperparameters, confirming the value of ImageNet/COCO "
        "transfer learning for this task."
    )
    doc.add_paragraph(
        "2. DINOv2 global-context fusion shows no measurable improvement on the primary metric "
        "(mAP@0.50: 0.4672 vs 0.4673, within noise), with small favorable deltas on secondary "
        "metrics (mAP@0.50:0.95, precision, recall) and a markedly more stable training curve "
        "(no overfitting collapse), attributable to its much smaller trainable parameter count."
    )
    doc.add_paragraph(
        "3. This negative-to-neutral fusion result is consistent with the domain-shift hypothesis "
        "flagged before experimentation: DINOv2 was pretrained on natural images, not radiographs, "
        "so its global embedding may carry limited chest-X-ray-specific signal. The finding is "
        "reported honestly rather than reframed, per project convention."
    )


def main() -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.size = Pt(10.5)

    add_title_page(doc)
    add_objective_and_dataset(doc)
    add_architecture(doc)
    add_training_configuration_table(doc)
    add_training_graphs(doc)
    add_comparison_graphs(doc)
    add_metrics_table(doc)
    add_qualitative_predictions(doc)
    add_conclusion(doc)

    doc.save(OUTPUT_PATH)
    print(f"report.docx saved to: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
